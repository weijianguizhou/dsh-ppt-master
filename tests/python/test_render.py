#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""lib/python/ppt_build.py 的渲染器测试（unittest，无第三方测试框架）。

覆盖：
  * probe() 返回结构与 --probe CLI 的哨兵输出协议；
  * 构造最小 spec 真实渲染 .pptx，断言文件可被 python-pptx 回读、页数正确；
  * 一组核心版式（kpi / chart / table / two-col / compare / quote）的冒烟渲染；
  * 预览 PNG 为「尽力而为」：生成成功就校验 PNG 魔数，环境不支持时允许降级为 warning。

python-pptx 缺失时相关用例整体 skip（CI 会先 pip install python-pptx pillow）。
注意：本文件不写死任何本机绝对路径，仓库根目录由 __file__ 反推。
"""

import contextlib
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(os.path.dirname(HERE))
SCRIPT = os.path.join(REPO_ROOT, 'lib', 'python', 'ppt_build.py')

try:  # python-pptx / Pillow 都是运行时依赖，缺失时允许 skip 而不是 error
    import pptx  # noqa: F401
    HAS_PPTX = True
except Exception:  # pragma: no cover - 取决于环境
    HAS_PPTX = False

try:
    import PIL  # noqa: F401
    HAS_PILLOW = True
except Exception:  # pragma: no cover - 取决于环境
    HAS_PILLOW = False

PNG_MAGIC = b'\x89PNG\r\n\x1a\n'


def load_renderer():
    """按文件路径导入渲染器模块（脚本不在包内，不能直接 import）。"""
    spec = importlib.util.spec_from_file_location('ppt_build_under_test', SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


RENDERER = load_renderer()


def run_cli(args, timeout=180):
    """调用 `python ppt_build.py ...`，返回 (returncode, payload, stdout, stderr)。"""
    env = dict(os.environ)
    env.setdefault('PYTHONUTF8', '1')
    env.setdefault('PYTHONIOENCODING', 'utf-8')
    proc = subprocess.run(
        [sys.executable, '-X', 'utf8', SCRIPT] + list(args),
        capture_output=True, text=True, timeout=timeout, env=env,
        encoding='utf-8', errors='replace',
    )
    payload = None
    for line in reversed((proc.stdout or '').splitlines()):
        line = line.strip()
        if line.startswith(RENDERER.SENTINEL):
            try:
                payload = json.loads(line[len(RENDERER.SENTINEL):])
            except Exception:
                payload = None
            break
    return proc.returncode, payload, proc.stdout, proc.stderr


def minimal_spec(out_path, preview=None):
    """三页最小 spec：cover + bullets + end，足以跑通渲染与回读。"""
    return {
        'template': {
            'id': 'indigo-corporate',
            'name': '靛蓝商务',
            'aspect': '16:9',
            'theme': {'mode': 'dark', 'bg': '#0A0F22', 'primary': '#5B7CFA', 'accent': '#FFC53D',
                      'text': '#F5F7FF', 'textDim': '#A7B2D6', 'surface': '#18224A', 'line': '#2B3765'},
            'fonts': {'title': {'latin': 'Arial', 'ea': 'Arial'},
                      'body': {'latin': 'Arial', 'ea': 'Arial'}},
            'geometry': {'marginX': 0.86, 'marginY': 0.58, 'bodySize': 16},
            'decor': {'style': 'glow', 'accentBar': True, 'coverPattern': 'orbs'},
        },
        'deck': {
            'title': '渲染器测试 deck',
            'subtitle': 'unittest smoke',
            'author': 'CI',
            'date': '2026-10',
            'slides': [
                {'layout': 'cover', 'title': '渲染器最小冒烟', 'subtitle': '三页 · 覆盖封面与结束页',
                 'notes': '开场一句话。'},
                {'layout': 'bullets', 'title': '只做三件事',
                 'bullets': ['第一件事：把结构理顺', '第二件事：把数据看清', '第三件事：把节奏跑对'],
                 'notes': '三件事按优先级排列。'},
                {'layout': 'end', 'title': '谢谢观看', 'contact': '测试团队',
                 'notes': '收尾。'},
            ],
        },
        'out': out_path,
    }


def rich_spec(out_path):
    """六种核心版式的冒烟 spec（chart / table / kpi / two-col / compare / quote）。"""
    spec = minimal_spec(out_path)
    spec['deck']['slides'] = [
        {'layout': 'cover', 'title': '核心版式冒烟', 'subtitle': 'python-pptx 真实渲染', 'notes': '开场。'},
        {'layout': 'kpi', 'title': '三个数字定义基线', 'notes': '三个指标。',
         'kpi': [{'value': '12.4', 'unit': '亿元', 'label': '全年营收', 'delta': '+18%', 'hint': '合并报表'},
                 {'value': '38', 'unit': '%', 'label': '直营占比', 'delta': '+17pt'},
                 {'value': '61', 'unit': '%', 'label': '复购率', 'delta': '+9pt'}]},
        {'layout': 'chart', 'title': '渠道结构在变化', 'insight': '直营贡献主要增量。', 'notes': '看两条柱子。',
         'chart': {'type': 'bar', 'categories': ['Q1', 'Q2', 'Q3', 'Q4'],
                   'series': [{'name': '直营', 'values': [1.2, 1.9, 2.6, 3.4]},
                              {'name': '代理', 'values': [2.1, 2.0, 2.2, 2.1]}]}},
        {'layout': 'table', 'title': '三条业务线效率对比', 'notes': '表格说明。',
         'table': {'head': ['业务线', '营收（亿元）', '毛利率'],
                   'rows': [['直营零售', '4.8', '58%'], ['企业服务', '3.6', '71%'], ['渠道代理', '4.0', '33%']]}},
        {'layout': 'two-col', 'title': '两种做法各有代价', 'notes': '左右对比说明。',
         'columns': [{'title': '继续补贴拉新', 'bullets': ['获客成本上涨', '复购率偏低']},
                     {'title': '押注复购与直营', 'bullets': ['存量毛利更高', '门店模型已跑通']}]},
        {'layout': 'compare', 'title': '只能选一条路', 'notes': '对比说明。',
         'left': {'title': '补贴拉新', 'items': ['成本高', '回报慢']},
         'right': {'title': '复购直营', 'items': ['毛利高', '可回本']}},
        {'layout': 'quote', 'title': '增长的本质是让同一批人买得更多',
         'quote': '把预算押在复购上，比再买一次流量更确定。', 'author': '战略部',
         'role': '年度汇报', 'notes': '题眼。'},
        {'layout': 'end', 'title': '谢谢观看', 'subtitle': '问答环节', 'notes': '收尾。'},
    ]
    return spec


class RendererContractTest(unittest.TestCase):
    """模块级契约：导出的入口、常量与布局注册表。"""

    def test_module_exports_expected_entrypoints(self):
        self.assertEqual(RENDERER.SENTINEL, '@@PPT_RESULT@@')
        for name in ('probe', 'build', 'main', 'resolve_layout'):
            self.assertTrue(callable(getattr(RENDERER, name, None)), '缺少入口：%s' % name)
        self.assertIsInstance(RENDERER.LAYOUTS, dict)
        self.assertGreaterEqual(len(RENDERER.LAYOUTS), 14)
        for layout in ('cover', 'bullets', 'kpi', 'chart', 'table', 'compare', 'end'):
            self.assertIn(layout, RENDERER.LAYOUTS)

    def test_resolve_layout_aliases_and_fallback(self):
        # 稳定别名
        self.assertEqual(RENDERER.resolve_layout({'layout': 'title'}), 'cover')
        self.assertEqual(RENDERER.resolve_layout({'layout': 'stats'}), 'kpi')
        self.assertEqual(RENDERER.resolve_layout({'layout': 'thanks'}), 'end')
        # 未知版式回退 bullets（与新版式开发解耦）
        self.assertEqual(RENDERER.resolve_layout({'layout': 'no-such-layout-xyz'}), 'bullets')
        self.assertEqual(RENDERER.resolve_layout({}), 'bullets')
        self.assertEqual(RENDERER.resolve_layout(None), 'bullets')

    def test_probe_returns_structured_environment(self):
        info = RENDERER.probe()
        self.assertIsInstance(info, dict)
        self.assertIn('ok', info)
        self.assertIn('python', info)
        if HAS_PPTX:
            self.assertTrue(info['ok'], info.get('error'))
            self.assertTrue(info.get('pythonPptx'))
        else:  # pragma: no cover - 取决于环境
            self.assertFalse(info['ok'])
            self.assertIn('error', info)

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_cli_probe_emits_sentinel_json(self):
        code, payload, stdout, stderr = run_cli(['--probe'])
        self.assertEqual(code, 0, stderr[-500:])
        self.assertIsInstance(payload, dict)
        self.assertTrue(payload['ok'], payload.get('error'))
        self.assertIn('pythonPptx', payload)
        self.assertIn(RENDERER.SENTINEL, stdout)


class RendererBuildTest(unittest.TestCase):
    """真实渲染：写 spec → 调 build/main → python-pptx 回读断言。"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix='ppt-master-test-')
        self.addCleanup(shutil.rmtree, self.tmpdir, True)

    def spec_path(self, spec, name='spec.json'):
        path = os.path.join(self.tmpdir, name)
        with open(path, 'w', encoding='utf-8') as handle:
            json.dump(spec, handle, ensure_ascii=False)
        return path

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_build_minimal_deck_and_read_back(self):
        out_path = os.path.join(self.tmpdir, 'minimal.pptx')
        spec = minimal_spec(out_path)
        result = RENDERER.build(spec)

        self.assertTrue(result.get('ok'), result.get('error'))
        self.assertTrue(os.path.exists(out_path), '渲染器未产出 pptx')
        self.assertGreater(os.path.getsize(out_path), 0)
        self.assertEqual(result['slides'], 3)
        self.assertEqual(result['layouts'], ['cover', 'bullets', 'end'])
        self.assertEqual(result.get('warnings'), [], '最小 spec 不应产生渲染警告')
        self.assertEqual(result.get('verified', {}).get('slides'), 3)
        self.assertGreater(result.get('verified', {}).get('shapes', 0), 0)

        reopened = pptx.Presentation(out_path)
        self.assertEqual(len(reopened.slides), 3)
        self.assertEqual(sum(len(slide.shapes) for slide in reopened.slides), result['verified']['shapes'])

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_cli_render_spec_emits_sentinel_and_valid_pptx(self):
        out_path = os.path.join(self.tmpdir, 'cli.pptx')
        spec_path = self.spec_path(minimal_spec(out_path))
        code, payload, stdout, stderr = run_cli([spec_path])
        self.assertEqual(code, 0, (stderr or stdout)[-800:])
        self.assertIsInstance(payload, dict)
        self.assertTrue(payload.get('ok'), payload.get('error'))
        self.assertTrue(os.path.exists(out_path))
        self.assertEqual(payload.get('slides'), 3)
        reopened = pptx.Presentation(out_path)
        self.assertEqual(len(reopened.slides), 3)

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_build_core_layouts_smoke(self):
        out_path = os.path.join(self.tmpdir, 'layouts.pptx')
        result = RENDERER.build(rich_spec(out_path))
        self.assertTrue(result.get('ok'), result.get('error'))
        self.assertEqual(result['slides'], 8)
        self.assertIn('chart', result['layouts'])
        self.assertIn('table', result['layouts'])
        self.assertEqual(result.get('warnings'), [], '核心版式渲染不应产生警告：%s' % result.get('warnings'))
        self.assertEqual(result.get('verified', {}).get('slides'), 8)

        # 图表与表格是原生对象：确认图表页真的有 graphicFrame / 表格形状
        reopened = pptx.Presentation(out_path)
        chart_shapes = [shape for slide in reopened.slides for shape in slide.shapes
                        if getattr(shape, 'has_chart', False)]
        table_shapes = [shape for slide in reopened.slides for shape in slide.shapes
                        if getattr(shape, 'has_table', False)]
        self.assertGreaterEqual(len(chart_shapes), 1, '应存在原生图表对象')
        self.assertGreaterEqual(len(table_shapes), 1, '应存在原生表格对象')
        self.assertEqual(len(chart_shapes[0].chart.plots[0].series), 2)

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_missing_slides_raises_clear_error(self):
        out_path = os.path.join(self.tmpdir, 'empty.pptx')
        spec = minimal_spec(out_path)
        spec['deck']['slides'] = []
        with self.assertRaisesRegex(Exception, 'deck\\.slides 为空|slides'):
            RENDERER.build(spec)
        self.assertFalse(os.path.exists(out_path))

    @unittest.skipUnless(HAS_PPTX and HAS_PILLOW, 'python-pptx / Pillow 未安装')
    def test_preview_png_is_best_effort(self):
        out_path = os.path.join(self.tmpdir, 'preview-deck.pptx')
        preview_path = os.path.join(self.tmpdir, 'preview.png')
        spec = minimal_spec(out_path, preview=preview_path)
        spec['preview'] = preview_path
        spec['previewGrid'] = 3
        spec['previewWidth'] = 320
        result = RENDERER.build(spec)
        self.assertTrue(result.get('ok'), result.get('error'))
        if os.path.exists(preview_path):
            with open(preview_path, 'rb') as handle:
                self.assertEqual(handle.read(8), PNG_MAGIC, '预览图应是合法 PNG')
            self.assertGreater(os.path.getsize(preview_path), 512)
        else:  # 字体/绘图环境不支持时允许降级为 warning
            self.assertTrue(
                any('预览' in str(warning) for warning in result.get('warnings', [])),
                '预览图未生成时应给出可解释的 warning：%s' % result.get('warnings'),
            )


class RendererCliTest(unittest.TestCase):
    """CLI 错误路径：给不存在的 spec 应优雅返回 ok=false，而不是抛栈退出。"""

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_missing_spec_path_returns_structured_error(self):
        code, payload, stdout, stderr = run_cli([])
        self.assertEqual(code, 1)
        self.assertIsInstance(payload, dict)
        self.assertFalse(payload.get('ok'))
        self.assertIn('spec', (payload.get('error') or '').lower())

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_nonexistent_spec_file_returns_structured_error(self):
        missing = os.path.join(tempfile.gettempdir(), 'ppt-master-no-such-spec-%d.json' % os.getpid())
        if os.path.exists(missing):
            os.remove(missing)
        code, payload, _stdout, _stderr = run_cli([missing])
        self.assertEqual(code, 1)
        self.assertIsInstance(payload, dict)
        self.assertFalse(payload.get('ok'))

    @unittest.skipUnless(HAS_PPTX, 'python-pptx 未安装')
    def test_cli_selftest_generates_sample_deck(self):
        out_dir = tempfile.mkdtemp(prefix='ppt-master-selftest-')
        self.addCleanup(shutil.rmtree, out_dir, True)
        code, payload, stdout, stderr = run_cli(['--selftest', out_dir])
        self.assertEqual(code, 0, (stderr or stdout)[-800:])
        self.assertIsInstance(payload, dict)
        self.assertTrue(payload.get('ok'), payload.get('error'))
        self.assertEqual(payload.get('slides'), 3)
        deck_path = os.path.join(out_dir, 'selftest.pptx')
        self.assertTrue(os.path.exists(deck_path), '--selftest 应产出 selftest.pptx')
        self.assertEqual(len(pptx.Presentation(deck_path).slides), 3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
